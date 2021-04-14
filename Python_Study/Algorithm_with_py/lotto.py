import random
import os
import openpyxl
from openpyxl import load_workbook

def is_In(Arr,a):
    if a in Arr : return True
    return False

def lotto_Make() :
    lotto = []
    for i in range(6):
        while(True):
            x = random.randrange(1, 46)
            if(not is_In(lotto,x)) : break

        lotto.append(x)


    return sorted(lotto) # 정렬해서 반환


def lotto_Rank(lotto,my_Lotto):
    # print("이번 당첨 번호 : ",end=" ")
    # print(lotto)
    # print("당신의 번호 : ",end=" ")
    # print(my_Lotto)
    k = 0
    for i in range(6):
        if lotto[i] in my_Lotto:
            k+=1

    if k == 6:
        my_Lotto[6]+=3000

    elif k == 5:
        indx = 0
        second = False
        while(indx<6): #7번은 점수임
            if lotto[6] == my_Lotto[indx] :
                second = True
                my_Lotto[6]+=600
            indx+=1
        # if lotto[6] in my_Lotto:  print("2등") # 보너스 번호가 있는 경우
        if(not second):
            my_Lotto[6] += 30
    elif k == 4: # 4개면 4등
        my_Lotto[6] += 10
    elif k == 3:
        my_Lotto[6] += 1 # 3개면 5등


rb = load_workbook("lotto.xlsx")
ws = rb['Sheet1']

row = ['b','c','d','e','f','g','h']
col = range(3,960)


# 1회차부터 현재까지의 모든 당첨번호를 배열로 가져옴
win_Num_Arr = []


for x in col:

    arr = []
    for i in row:
        arr.append(ws[i+str(x)].value)
    win_Num_Arr.append(arr)




lotto_Arr = []

for i in range(8100000): # 로또 번호생성, 배열에 추가
    a = lotto_Make() #로또 번호생성
    a.append(0) #처음에 평가하기 전, 점수
    lotto_Arr.append(a) #배열에 추가



# win_Num = [4,15,24,35,36,40,1]

wb = openpyxl.Workbook()
sheet1 = wb.active

for i in lotto_Arr:
    for x in win_Num_Arr:
        lotto_Rank(x,i)
    sheet1.append([i[0],i[1],i[2],i[3],i[4],i[5],i[6]])

wb.save("lotto_1.xlsx") # 엑셀에 저장


